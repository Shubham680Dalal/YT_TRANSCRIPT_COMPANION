"""Per-video watch progress, kept in memory and persisted to an Excel sheet.

The browser sends a heartbeat every few seconds while a video plays. Those only touch
the in-memory copy; the sheet is written on pause / end / tab close, every few minutes
while playing (see backend/app.py), and once more when the server shuts down (Ctrl+C).

The sheet is meant to be human-readable and hand-editable: change "Left off at" in
Excel, save, and the next read picks it up (we reload whenever the file's mtime changes).
"""
import atexit
import datetime as dt
import os
import threading
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from backend.utils.config_loader import get_config
from backend.utils.logger import get_logger

log = get_logger(__name__)

PROJECT_ROOT = Path(__file__).parent.parent.parent

HEADERS = ["Video ID", "Title", "Link", "Left off at", "Duration", "% watched", "Status", "Last watched"]
COLUMN_WIDTHS = [14, 60, 45, 12, 12, 11, 12, 18]
STATUS_IN_PROGRESS = "in progress"
STATUS_FINISHED = "finished"


def format_timestamp(seconds: float) -> str:
    s = int(seconds)
    h, m, sec = s // 3600, (s % 3600) // 60, s % 60
    return f"{h}:{m:02d}:{sec:02d}" if h else f"{m}:{sec:02d}"


def parse_timestamp(value) -> float | None:
    """Accepts what we write ("1:02:05") and what Excel turns a hand-typed time into."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, dt.timedelta):
        return value.total_seconds()
    if isinstance(value, dt.datetime):  # Excel stores times past 24h as datetimes from 1899-12-30
        return (value - dt.datetime(1899, 12, 30)).total_seconds()
    if isinstance(value, (int, float)):
        return float(value)
    try:
        seconds = 0.0
        for part in str(value).strip().split(":"):
            seconds = seconds * 60 + float(part)
        return seconds
    except ValueError:
        return None


def watch_url(video_id: str) -> str:
    return f"https://www.youtube.com/watch?v={video_id}"


class ProgressStore:
    def __init__(self, path: Path, min_seconds: float = 5, end_margin_seconds: float = 10):
        self.path = Path(path)
        self.min_seconds = min_seconds
        self.end_margin_seconds = end_margin_seconds
        self._records: dict[str, dict] = {}
        self._heartbeats: dict[str, float] = {}  # video_id -> time.monotonic() of last update
        self._dirty = False
        self._file_mtime: float | None = None
        self._lock = threading.RLock()
        self._load()

    # ---- disk ----

    def _load(self) -> None:
        if not self.path.exists():
            self._file_mtime = None
            return
        try:
            wb = load_workbook(self.path)
        except Exception as exc:
            log.error("progress_load_failed", path=str(self.path), error=str(exc))
            return
        ws = wb.active
        loaded = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = list(row) + [None] * (len(HEADERS) - len(row))
            video_id, title, _link, left_off, duration, _pct, status, updated = row[:len(HEADERS)]
            if not video_id:
                continue
            loaded[str(video_id).strip()] = {
                "video_id": str(video_id).strip(),
                "title": title or "",
                "position": parse_timestamp(left_off) or 0.0,
                "duration": parse_timestamp(duration) or 0.0,
                "status": status or STATUS_IN_PROGRESS,
                "updated_at": updated if isinstance(updated, dt.datetime) else dt.datetime.now(),
            }
        self._records = loaded
        self._file_mtime = self.path.stat().st_mtime
        log.info("progress_loaded", path=str(self.path), videos=len(loaded))

    def _reload_if_changed_on_disk(self) -> None:
        """Pick up hand edits made in Excel, keeping any newer unsaved in-memory changes."""
        if not self.path.exists():
            return
        mtime = self.path.stat().st_mtime
        if self._file_mtime is not None and mtime == self._file_mtime:
            return
        unsaved = dict(self._records) if self._dirty else {}
        self._load()
        for video_id, record in unsaved.items():
            on_disk = self._records.get(video_id)
            if on_disk is None or record["updated_at"] > on_disk["updated_at"]:
                self._records[video_id] = record
        if unsaved:
            log.info("progress_reloaded_after_external_edit", path=str(self.path))

    def flush(self) -> bool:
        """Write the sheet if anything changed. Returns True if the file was written."""
        with self._lock:
            if not self._dirty:
                return False
            self._reload_if_changed_on_disk()
            wb = Workbook()
            ws = wb.active
            ws.title = "Watch progress"
            ws.append(HEADERS)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            for i, width in enumerate(COLUMN_WIDTHS):
                ws.column_dimensions[chr(ord("A") + i)].width = width

            for rec in sorted(self._records.values(), key=lambda r: r["updated_at"], reverse=True):
                pct = round(100 * rec["position"] / rec["duration"]) if rec["duration"] else ""
                if rec["status"] == STATUS_FINISHED:
                    pct = 100
                ws.append([
                    rec["video_id"],
                    rec["title"],
                    watch_url(rec["video_id"]),
                    format_timestamp(rec["position"]),
                    format_timestamp(rec["duration"]) if rec["duration"] else "",
                    pct,
                    rec["status"],
                    rec["updated_at"],
                ])
                row = ws.max_row
                ws.cell(row=row, column=3).hyperlink = watch_url(rec["video_id"])
                ws.cell(row=row, column=3).style = "Hyperlink"
                ws.cell(row=row, column=8).number_format = "yyyy-mm-dd hh:mm"

            self.path.parent.mkdir(parents=True, exist_ok=True)
            tmp_path = self.path.with_name(f".{self.path.stem}.tmp.xlsx")
            try:
                wb.save(tmp_path)
                os.replace(tmp_path, self.path)
            except PermissionError:
                # Usually: the sheet is open in Excel, which locks it on Windows.
                log.warning("progress_file_locked", path=str(self.path),
                            hint="close the Excel file - changes are kept and saved on the next write")
                tmp_path.unlink(missing_ok=True)
                return False
            self._file_mtime = self.path.stat().st_mtime
            self._dirty = False
            log.info("progress_saved", path=str(self.path), videos=len(self._records))
            return True

    # ---- reads / writes ----

    def get(self, video_id: str) -> dict | None:
        with self._lock:
            self._reload_if_changed_on_disk()
            record = self._records.get(video_id)
            return dict(record) if record else None

    def list_recent(self, limit: int = 20) -> list[dict]:
        with self._lock:
            self._reload_if_changed_on_disk()
            ordered = sorted(self._records.values(), key=lambda r: r["updated_at"], reverse=True)
            return [dict(r) for r in ordered[:limit]]

    def update(self, video_id: str, position: float, duration: float = 0,
               title: str | None = None, state: str = "playing") -> dict:
        with self._lock:
            self._reload_if_changed_on_disk()
            record = self._records.get(video_id) or {
                "video_id": video_id, "title": "", "position": 0.0,
                "duration": 0.0, "status": STATUS_IN_PROGRESS,
            }
            if title:
                record["title"] = title
            if duration > 0:
                record["duration"] = duration

            near_end = record["duration"] > 0 and position >= record["duration"] - self.end_margin_seconds
            if state == "ended" or near_end:
                record["status"] = STATUS_FINISHED
                record["position"] = 0.0
            elif position >= self.min_seconds:
                record["status"] = STATUS_IN_PROGRESS
                record["position"] = position
            # else: a 0:00-ish report (fresh page, player not started yet) - never let it
            # overwrite a real saved position.

            record["updated_at"] = dt.datetime.now().replace(microsecond=0)
            self._records[video_id] = record
            self._heartbeats[video_id] = time.monotonic()
            self._dirty = True
            return dict(record)

    def recently_active(self, within_seconds: float = 30) -> list[dict]:
        cutoff = time.monotonic() - within_seconds
        with self._lock:
            return [dict(self._records[v]) for v, t in self._heartbeats.items()
                    if t >= cutoff and v in self._records]


_store: ProgressStore | None = None


def get_store() -> ProgressStore:
    global _store
    if _store is None:
        cfg = get_config()["progress"]
        _store = ProgressStore(
            PROJECT_ROOT / cfg["file"],
            min_seconds=cfg["min_seconds"],
            end_margin_seconds=cfg["end_margin_seconds"],
        )
        atexit.register(_store.flush)  # backup in case shutdown skips the lifespan hook
    return _store


def flush_on_shutdown() -> None:
    store = get_store()
    for rec in store.recently_active():
        log.info("progress_saved_on_shutdown", title=rec["title"] or rec["video_id"],
                 left_off_at=format_timestamp(rec["position"]), status=rec["status"])
    store.flush()
