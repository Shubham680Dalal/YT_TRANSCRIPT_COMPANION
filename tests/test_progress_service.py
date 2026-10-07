import datetime as dt
import os

import pytest
from openpyxl import load_workbook

from backend.services import progress_service
from backend.services.progress_service import ProgressStore, parse_timestamp

VID = "aircAruvnKk"
OTHER = "2T86xAtR6Fo"


@pytest.fixture
def sheet(tmp_path):
    return tmp_path / "data" / "watch_progress.xlsx"


def test_update_and_read_back(sheet):
    store = ProgressStore(sheet)
    store.update(VID, 125.4, duration=600, title="Neural nets", state="paused")

    rec = store.get(VID)
    assert rec["position"] == 125.4
    assert rec["title"] == "Neural nets"
    assert rec["status"] == "in progress"


def test_near_zero_report_never_overwrites_saved_spot(sheet):
    store = ProgressStore(sheet)
    store.update(VID, 300, duration=600, state="paused")
    store.update(VID, 0, duration=600, state="closed")  # e.g. refresh before pressing play

    assert store.get(VID)["position"] == 300


def test_ended_or_near_end_counts_as_finished(sheet):
    store = ProgressStore(sheet)
    store.update(VID, 595, duration=600, state="playing")
    store.update(OTHER, 50, duration=600, state="ended")

    assert store.get(VID)["status"] == "finished"
    assert store.get(VID)["position"] == 0
    assert store.get(OTHER)["status"] == "finished"


def test_flush_writes_readable_sheet_and_new_store_loads_it(sheet):
    store = ProgressStore(sheet)
    store.update(VID, 3725, duration=7200, title="Long one", state="paused")
    assert store.flush() is True
    assert store.flush() is False  # nothing changed since

    ws = load_workbook(sheet).active
    assert [c.value for c in ws[1]][:4] == ["Video ID", "Title", "Link", "Left off at"]
    assert ws["A2"].value == VID
    assert ws["D2"].value == "1:02:05"
    assert ws["C2"].hyperlink.target == f"https://www.youtube.com/watch?v={VID}"

    reloaded = ProgressStore(sheet)
    assert reloaded.get(VID)["position"] == 3725
    assert reloaded.get(VID)["title"] == "Long one"


def test_hand_edit_in_excel_is_picked_up(sheet):
    store = ProgressStore(sheet)
    store.update(VID, 100, duration=600, state="paused")
    store.flush()

    wb = load_workbook(sheet)
    wb.active["D2"] = "7:30"
    wb.save(sheet)
    st = os.stat(sheet)
    os.utime(sheet, (st.st_atime, st.st_mtime + 5))  # make sure the mtime visibly changes

    assert store.get(VID)["position"] == 450


def test_locked_file_keeps_changes_for_next_flush(sheet, monkeypatch):
    store = ProgressStore(sheet)
    store.update(VID, 100, state="paused")

    def locked(*args, **kwargs):
        raise PermissionError("file is open in Excel")

    monkeypatch.setattr(progress_service.os, "replace", locked)
    assert store.flush() is False
    assert not sheet.exists()

    monkeypatch.undo()
    assert store.flush() is True
    assert ProgressStore(sheet).get(VID)["position"] == 100


def test_list_recent_is_newest_first(sheet):
    store = ProgressStore(sheet)
    store.update(VID, 10, state="paused")
    store.update(OTHER, 20, state="paused")
    store._records[VID]["updated_at"] -= dt.timedelta(minutes=5)

    assert [r["video_id"] for r in store.list_recent()] == [OTHER, VID]


@pytest.mark.parametrize("value, expected", [
    ("1:02:05", 3725),
    ("7:30", 450),
    (dt.time(0, 7, 30), 450),
    (dt.timedelta(hours=1, seconds=5), 3605),
    (90, 90),
    ("", None),
    ("abc", None),
])
def test_parse_timestamp(value, expected):
    assert parse_timestamp(value) == expected
